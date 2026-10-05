"""Builds business/knock-unit-economics.xlsx (editable inputs, formulas). Run: python3 business/build_unit_economics.py"""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

wb = Workbook()
inp = wb.active; inp.title = "Inputs"
blue = Font(color="1F4E9E"); bold = Font(bold=True); fill = PatternFill("solid", fgColor="FFF4CC")
rows = [
    ("General", None, None),
    ("USD to INR", 88, ""),
    ("Payment gateway fee % (Razorpay 2% + 18% GST on fee)", 0.0236, ""),
    ("Charge GST to client? (1 yes, 0 no; needed above Rs 20 lakh turnover)", 0, ""),
    ("GST rate", 0.18, ""),
    ("Founder or operator cost per hour (Rs)", 500, "what your time is worth"),
    ("Monthly fixed costs (Rs): Vercel Pro, domain, misc", 2500, ""),
    ("Clients per month (to spread fixed costs)", 5, ""),
    ("Model prices (USD per million tokens)", None, None),
    ("Opus 5.5 input", 4, ""), ("Opus 5.5 output", 20, ""),
    ("Sonnet 5.5 input", 2, ""), ("Sonnet 5.5 output", 10, ""),
    ("Haiku 4.5 input", 1, ""), ("Haiku 4.5 output", 5, ""),
    ("Web search, USD per search", 0.01, ""),
    ("Prompt caching saving on input (0 to 0.6)", 0.3, "repeated system prompts"),
    ("Per-prospect work (output includes thinking). Model: 1 Opus, 2 Sonnet, 3 Haiku", None, None),
    ("Classify: input tokens", 2000, ""), ("Classify: output tokens", 300, ""), ("Classify: model", 3, ""),
    ("Enrich: searches", 2, ""), ("Enrich: input tokens", 6000, ""), ("Enrich: output tokens", 800, ""), ("Enrich: model", 2, ""),
    ("Hooks: input tokens", 6000, ""), ("Hooks: output tokens", 1500, ""), ("Hooks: model", 1, ""),
    ("Writers: count", 3, ""), ("Writer: input tokens each", 3000, ""), ("Writer: output tokens each", 600, ""), ("Writers: model", 2, ""),
    ("Judge: input tokens", 4000, ""), ("Judge: output tokens", 800, ""), ("Judge: model", 1, ""),
    ("Discovery searches per prospect (amortised)", 1, ""),
    ("Conversation", None, None),
    ("Share of prospects that get a follow-up", 0.7, ""),
    ("Follow-up: input tokens", 3000, ""), ("Follow-up: output tokens", 500, ""), ("Follow-up: model", 1, ""),
    ("Reply rate", 0.3, "replies per prospect"),
    ("Exchanges per reply", 2, ""),
    ("Reply draft: input tokens", 5000, ""), ("Reply draft: output tokens", 800, ""), ("Reply draft: model", 1, ""),
    ("Per client setup (USD): profile extraction + learning runs", 0.5, ""),
]
inp["A1"] = "Knock unit economics: change the yellow cells"; inp["A1"].font = Font(bold=True, size=14)
N = {}; r = 3
for label, val, note in rows:
    if val is None:
        inp.cell(r, 1, label).font = bold
    else:
        inp.cell(r, 1, label); c = inp.cell(r, 2, val); c.font = blue; c.fill = fill; inp.cell(r, 3, note)
        N[label] = f"Inputs!$B${r}"
    r += 1
inp.column_dimensions["A"].width = 72; inp.column_dimensions["B"].width = 14; inp.column_dimensions["C"].width = 28

def price(m, io):
    if io == "in":
        return f"CHOOSE({m},{N['Opus 5.5 input']},{N['Sonnet 5.5 input']},{N['Haiku 4.5 input']})*(1-{N['Prompt caching saving on input (0 to 0.6)']})"
    return f"CHOOSE({m},{N['Opus 5.5 output']},{N['Sonnet 5.5 output']},{N['Haiku 4.5 output']})"

def step(i, o, m, mult="1"):
    return f"({mult})*({N[i]}*{price(N[m], 'in')}+{N[o]}*{price(N[m], 'out')})/1000000"

cost = wb.create_sheet("Cost per prospect")
items = [
    ("Classify", step("Classify: input tokens", "Classify: output tokens", "Classify: model")),
    ("Enrich (model)", step("Enrich: input tokens", "Enrich: output tokens", "Enrich: model")),
    ("Enrich (searches)", f"{N['Enrich: searches']}*{N['Web search, USD per search']}"),
    ("Discovery searches", f"{N['Discovery searches per prospect (amortised)']}*{N['Web search, USD per search']}"),
    ("Hooks", step("Hooks: input tokens", "Hooks: output tokens", "Hooks: model")),
    ("Writers", step("Writer: input tokens each", "Writer: output tokens each", "Writers: model", N['Writers: count'])),
    ("Judge", step("Judge: input tokens", "Judge: output tokens", "Judge: model")),
    ("Follow-ups", step("Follow-up: input tokens", "Follow-up: output tokens", "Follow-up: model", N['Share of prospects that get a follow-up'])),
    ("Reply drafts", step("Reply draft: input tokens", "Reply draft: output tokens", "Reply draft: model", f"{N['Reply rate']}*{N['Exchanges per reply']}")),
]
cost["A1"] = "Claude cost per prospect"; cost["A1"].font = Font(bold=True, size=14)
for c, t in (("A3", "Step"), ("B3", "USD"), ("C3", "INR")):
    cost[c] = t; cost[c].font = bold
r = 4
for lab, f in items:
    cost.cell(r, 1, lab); cost.cell(r, 2, "=" + f).number_format = "0.0000"
    cost.cell(r, 3, f"=B{r}*{N['USD to INR']}").number_format = "0.00"; r += 1
cost.cell(r, 1, "Total per prospect").font = bold
cost.cell(r, 2, f"=SUM(B4:B{r-1})").number_format = "0.0000"
cost.cell(r, 3, f"=SUM(C4:C{r-1})").number_format = "0.00"
TOT = f"'Cost per prospect'!$C${r}"
for col, w in (("A", 28), ("B", 12), ("C", 12)):
    cost.column_dimensions[col].width = w

pk = wb.create_sheet("Packages", 0)
hdr = ["Package", "Price (Rs)", "People", "Founder hours", "Tool cost (Rs, e.g. LGM)", "Claude cost (Rs)", "Setup (Rs)",
       "Gateway fee (Rs)", "GST (Rs)", "Fixed cost share (Rs)", "Founder time (Rs)", "Total cost (Rs)", "Profit (Rs)",
       "Margin %", "Cash margin % (excl. founder time)"]
pk["A1"] = "Cost, price and margin per client: change the yellow cells here and on Inputs"; pk["A1"].font = Font(bold=True, size=14)
for i, h in enumerate(hdr, 1):
    c = pk.cell(3, i, h); c.font = bold; c.alignment = Alignment(wrap_text=True)
packs = [("Sprint, 7 days", 2499, 60, 0.5, 0), ("Standard, 14 days", 3999, 120, 0.75, 0), ("Full, 21 days", 5499, 180, 1, 0),
         ("Done-for-you (later, not offered now)", 19999, 180, 16, 6160)]
for j, (n, p, ppl, h, tool) in enumerate(packs):
    r = 4 + j
    pk.cell(r, 1, n)
    for col, v in ((2, p), (3, ppl), (4, h), (5, tool)):
        c = pk.cell(r, col, v); c.font = blue; c.fill = fill
    pk.cell(r, 6, f"=C{r}*{TOT}")
    pk.cell(r, 7, f"={N['Per client setup (USD): profile extraction + learning runs']}*{N['USD to INR']}")
    pk.cell(r, 8, f"=B{r}*{N['Payment gateway fee % (Razorpay 2% + 18% GST on fee)']}")
    pk.cell(r, 9, f"=B{r}-B{r}/(1+{N['GST rate']}*{N['Charge GST to client? (1 yes, 0 no; needed above Rs 20 lakh turnover)']})")
    pk.cell(r, 10, f"={N['Monthly fixed costs (Rs): Vercel Pro, domain, misc']}/{N['Clients per month (to spread fixed costs)']}")
    pk.cell(r, 11, f"=D{r}*{N['Founder or operator cost per hour (Rs)']}")
    pk.cell(r, 12, f"=SUM(E{r}:K{r})")
    pk.cell(r, 13, f"=B{r}-L{r}")
    pk.cell(r, 14, f"=M{r}/B{r}").number_format = "0%"
    pk.cell(r, 15, f"=(M{r}+K{r})/B{r}").number_format = "0%"
    for col in range(6, 14):
        pk.cell(r, col).number_format = "#,##0"
for i in range(1, 16):
    pk.column_dimensions[get_column_letter(i)].width = 14
pk.column_dimensions["A"].width = 26
pk["A11"] = ("Notes: Claude cost uses the step assumptions on Inputs (estimates until measured with the cost command). "
             "Done-for-you tool cost = one La Growth Machine Basic identity month (about $70). Founder hours = support and spot checks only (automation runs as scheduled Claude Code tasks; the client approves and sends).")
wb.save(Path(__file__).with_name("knock-unit-economics.xlsx"))
